"""Real-data onboarding: turn common platform exports into the weekly data contract.

Mappers exist for Meta Ads Manager, Google Ads and Shopify order exports. Each returns a
weekly frame (weeks start on Monday) that ``build_contract_frame`` joins into the contract
checked by ``mixlab.validate``. ``anonymize`` scales all money by one secret factor so a
brand can share data without revealing its true revenue or spend.
"""

import hashlib
import io
import re
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from mixlab import config

MONEY_JUNK = re.compile(r"[^\d.\-]")
ISO_DATE = re.compile(r"^\d{4}-\d{2}-\d{2}")


class OnboardingError(ValueError):
    """An export could not be mapped; the message says what to fix."""


# --- Parsing --------------------------------------------------------------------------------


def parse_money(values: pd.Series) -> pd.Series:
    """Convert money written as text ("₹1,23,456.50", "Rs 900", "(120)") to numbers.

    Currency symbols, codes, spaces and thousands separators (Indian or Western grouping)
    are removed. Parentheses mean negative. Blank cells become zero.
    """
    if pd.api.types.is_numeric_dtype(values):
        return values.fillna(0.0).astype(float)
    text = values.fillna("").astype(str).str.strip()
    wordy = (text != "") & (text != "-") & ~text.str.contains(r"\d")
    if wordy.any():
        raise OnboardingError(f"Could not read '{text[wordy].iloc[0]}' as an amount of money.")
    negative = text.str.startswith("(") & text.str.endswith(")")
    cleaned = text.str.replace(MONEY_JUNK, "", regex=True).replace({"": "0", "-": "0"})
    numbers = pd.to_numeric(cleaned, errors="coerce")
    if numbers.isna().any():
        bad = text[numbers.isna()].iloc[0]
        raise OnboardingError(f"Could not read '{bad}' as an amount of money.")
    return numbers.where(~negative, -numbers.abs()).astype(float)


def parse_dates(values: pd.Series, dayfirst: bool | None = None) -> pd.Series:
    """Parse dates in the formats platforms export, returning timezone-free dates.

    ISO dates (2024-03-05, with or without a time and UTC offset) are read directly. For
    slash or dash dates the day/month order is inferred: if any first number exceeds 12 the
    file is day-first. Pass ``dayfirst`` to override when every date is ambiguous.
    """
    text = values.astype(str).str.strip()
    if text.str.match(ISO_DATE).all():
        # Keep the local calendar date: drop the time and any UTC offset.
        return pd.to_datetime(text.str.slice(0, 10), format="%Y-%m-%d")
    if dayfirst is None:
        first_number = pd.to_numeric(text.str.extract(r"^(\d{1,2})[/\-.]")[0], errors="coerce")
        dayfirst = bool((first_number > 12).any())
    parsed = pd.to_datetime(text, format="mixed", dayfirst=dayfirst, errors="coerce")
    if parsed.isna().any():
        raise OnboardingError(f"Could not read '{text[parsed.isna()].iloc[0]}' as a date.")
    return parsed.dt.tz_localize(None).dt.normalize()


def find_column(df: pd.DataFrame, candidates: tuple[str, ...], what: str) -> str:
    """Return the first column whose name starts with one of ``candidates`` (case-insensitive)."""
    for candidate in candidates:
        for column in df.columns:
            if str(column).strip().lower().startswith(candidate):
                return column
    raise OnboardingError(
        f"No {what} column found. Expected a column starting with one of {list(candidates)}; "
        f"the file has {list(df.columns)}."
    )


def read_export(path: Path, header_keyword: str) -> pd.DataFrame:
    """Read a CSV that may have title lines above the header (as Google Ads reports do).

    The header is the first line containing ``header_keyword``. Trailing "Total" rows are
    dropped.
    """
    lines = Path(path).read_text(encoding="utf-8-sig").splitlines()
    for index, line in enumerate(lines):
        if header_keyword.lower() in line.lower():
            frame = pd.read_csv(io.StringIO("\n".join(lines[index:])))
            first = frame.iloc[:, 0].astype(str).str.strip().str.lower()
            return frame.loc[~first.str.startswith("total")].reset_index(drop=True)
    raise OnboardingError(f"No header row containing '{header_keyword}' found in {path}.")


def to_weekly(daily: pd.DataFrame) -> pd.DataFrame:
    """Sum daily (or finer) rows into weeks starting on Monday.

    Args:
        daily: Frame with a ``date`` column and numeric columns to sum.

    Returns:
        One row per week from the first to the last week present; weeks with no rows are 0.

    """
    week = daily[config.DATE_COL].dt.to_period(config.WEEK_PERIOD)
    weekly = daily.drop(columns=config.DATE_COL).groupby(week.dt.start_time).sum()
    full = pd.date_range(weekly.index.min(), weekly.index.max(), freq=f"{config.DAYS_PER_WEEK}D")
    return weekly.reindex(full, fill_value=0.0).rename_axis(config.DATE_COL).reset_index()


# --- Platform mappers -----------------------------------------------------------------------


def map_meta_ads(df: pd.DataFrame, channel: str = "meta_ads") -> pd.DataFrame:
    """Map a Meta Ads Manager export to weekly ``spend_<channel>``.

    Expects a daily breakdown ("Day" or "Reporting starts") and an "Amount spent (...)"
    column. Multiple rows per day (campaigns, ad sets) are summed.
    """
    date_column = find_column(df, ("day", "reporting starts", "date"), "date")
    spend_column = find_column(df, ("amount spent",), "amount spent")
    daily = pd.DataFrame(
        {
            config.DATE_COL: parse_dates(df[date_column]),
            f"{config.SPEND_PREFIX}{channel}": parse_money(df[spend_column]),
        }
    )
    return to_weekly(daily)


def map_google_ads(df: pd.DataFrame, split_by_campaign_type: bool = True) -> pd.DataFrame:
    """Map a Google Ads report to weekly spend, one column per campaign type.

    Expects "Day" (or "Date") and "Cost". If a "Campaign type" column is present and
    ``split_by_campaign_type`` is true, Search, Video (YouTube), Shopping, Performance Max and
    Display become separate channels; otherwise everything is ``spend_google_ads``.
    """
    date_column = find_column(df, ("day", "date"), "date")
    cost_column = find_column(df, ("cost",), "cost")
    channel = pd.Series("google_ads", index=df.index)
    type_columns = [c for c in df.columns if str(c).strip().lower() == "campaign type"]
    if split_by_campaign_type and type_columns:
        kinds = df[type_columns[0]].astype(str).str.strip().str.lower()
        channel = kinds.map(config.GOOGLE_CAMPAIGN_CHANNELS).fillna("google_other")
    long = pd.DataFrame(
        {
            config.DATE_COL: parse_dates(df[date_column]),
            "channel": config.SPEND_PREFIX + channel,
            "cost": parse_money(df[cost_column]),
        }
    )
    daily = long.pivot_table(
        index=config.DATE_COL, columns="channel", values="cost", aggfunc="sum", fill_value=0.0
    )
    daily.columns.name = None
    return to_weekly(daily.reset_index())


def map_shopify_orders(df: pd.DataFrame) -> pd.DataFrame:
    """Map a Shopify orders export to weekly ``revenue``.

    The export repeats an order on one row per line item with the total on the first row
    only, so orders are de-duplicated by "Name". Cancelled, voided and fully refunded orders
    are excluded, and partial refunds are subtracted. Orders are dated by "Created at" in the
    store's own time zone.

    Raises:
        OnboardingError: If the export mixes currencies.

    """
    required = {"Name", "Created at", "Total"}
    if missing := required - set(df.columns):
        raise OnboardingError(f"Shopify export is missing these columns: {sorted(missing)}.")
    orders = df.dropna(subset=["Total"]).drop_duplicates(subset="Name", keep="first").copy()
    if "Currency" in orders and orders["Currency"].nunique() > 1:
        currencies = sorted(orders["Currency"].dropna().unique())
        raise OnboardingError(
            f"The export mixes currencies {currencies}. Convert to INR or export one currency."
        )
    if "Cancelled at" in orders:
        orders = orders[orders["Cancelled at"].isna()]
    if "Financial Status" in orders:
        status = orders["Financial Status"].astype(str).str.lower()
        orders = orders[~status.isin(config.SHOPIFY_EXCLUDED_STATUSES)]
    revenue = parse_money(orders["Total"])
    if "Refunded Amount" in orders:
        revenue = revenue - parse_money(orders["Refunded Amount"])
    daily = pd.DataFrame(
        {config.DATE_COL: parse_dates(orders["Created at"]), config.TARGET_COL: revenue}
    )
    return to_weekly(daily)


def map_generic(
    df: pd.DataFrame,
    date_column: str,
    revenue_column: str,
    spend_columns: dict[str, str],
    control_columns: dict[str, str] | None = None,
) -> pd.DataFrame:
    """Rename an already-weekly dataset into the data contract.

    Args:
        df: Source frame, one row per week.
        date_column: Name of the week column.
        revenue_column: Name of the target column.
        spend_columns: Source column -> channel name (the ``spend_`` prefix is added).
        control_columns: Source column -> contract column name, for controls to keep.

    """
    renames = {date_column: config.DATE_COL, revenue_column: config.TARGET_COL}
    renames |= {source: f"{config.SPEND_PREFIX}{name}" for source, name in spend_columns.items()}
    renames |= control_columns or {}
    frame = df[list(renames)].rename(columns=renames)
    frame[config.DATE_COL] = parse_dates(frame[config.DATE_COL])
    return frame


def build_contract_frame(
    revenue: pd.DataFrame, spend: list[pd.DataFrame], controls: pd.DataFrame | None = None
) -> pd.DataFrame:
    """Join weekly revenue, spend and controls into one contract frame.

    The weeks covered are those with revenue. A channel with no spend recorded in a week is
    set to 0 for that week (the platform had nothing to report); missing controls are left
    empty so the validator flags them.
    """
    frame = revenue[[config.DATE_COL, config.TARGET_COL]].copy()
    for weekly in spend:
        frame = frame.merge(weekly, on=config.DATE_COL, how="left")
    spend_names = [c for c in frame.columns if c.startswith(config.SPEND_PREFIX)]
    frame[spend_names] = frame[spend_names].fillna(0.0)
    if controls is not None:
        frame = frame.merge(controls, on=config.DATE_COL, how="left")
    return frame.sort_values(config.DATE_COL).reset_index(drop=True)


# --- Anonymization --------------------------------------------------------------------------


def factor_from_secret(secret: str) -> float:
    """Return the scale factor derived from a passphrase.

    The same passphrase always gives the same factor, so the brand can reverse it later; it
    cannot be recovered from the scaled data without knowing a true figure.
    """
    if not secret:
        raise OnboardingError("The anonymization secret cannot be empty.")
    low, high = config.ANONYMIZE_FACTOR_RANGE
    fraction = int(hashlib.sha256(secret.encode()).hexdigest()[:12], 16) / 16**12
    return low + fraction * (high - low)


def money_columns(df: pd.DataFrame) -> list[str]:
    """Return the columns that hold money: revenue and every spend column."""
    return [c for c in df.columns if c == config.TARGET_COL or c.startswith(config.SPEND_PREFIX)]


def anonymize(df: pd.DataFrame, factor: float) -> pd.DataFrame:
    """Scale revenue and all spend by one factor.

    Because every money column is scaled together, ROI, channel shares, response shapes and
    timing are unchanged, so the model's conclusions carry over; only absolute rupee levels
    are hidden. Dates, seasonality and the channel mix remain visible.
    """
    if factor <= 0:
        raise OnboardingError("The anonymization factor must be positive.")
    scaled = df.copy()
    columns = money_columns(df)
    scaled[columns] = scaled[columns] * factor
    return scaled


def deanonymize(df: pd.DataFrame, factor: float) -> pd.DataFrame:
    """Reverse ``anonymize`` with the same factor."""
    return anonymize(df, 1.0 / factor)


# --- Excel template -------------------------------------------------------------------------

TEMPLATE_COLUMNS: dict[str, str] = {
    "date": "Monday that starts the week, as YYYY-MM-DD. One row per week, no gaps.",
    "revenue": "Net revenue for the week in INR (after discounts and refunds, before tax).",
    "spend_meta_ads": "Meta (Facebook + Instagram) ad spend for the week in INR.",
    "spend_google_search": "Google Search ad spend for the week in INR.",
    "spend_youtube": "YouTube / Google Video ad spend for the week in INR.",
    "spend_influencers": "Influencer fees for the week in INR, dated when the content went live.",
    "promo_flag": "1 if a sitewide promotion ran that week, otherwise 0.",
    "holiday_diwali": "1 in the week of Diwali, otherwise 0. Add one holiday_<name> column each.",
    "price_index": "Average selling price relative to the first week (first week = 1.00).",
}
TEMPLATE_INSTRUCTIONS: tuple[str, ...] = (
    "MixLab data template",
    "",
    "1. Fill in the 'Data' sheet: one row per week, weeks starting on Monday.",
    "2. Keep the column names exactly as they are. Add one spend_<channel> column for every "
    "channel you spend on, and delete example channels you do not use.",
    "3. Provide at least 104 weeks (two years). Three years is better.",
    "4. All money in INR, as plain numbers (no currency symbols or commas).",
    "5. A week with no spend on a channel is 0, not blank.",
    "6. Do not skip weeks, and do not add totals or notes below the data.",
    "7. The 'Columns' sheet explains each column. The example rows show the format; replace "
    "them with your own.",
    "",
    "When done, upload the file in the app's Upload page (as CSV) or send it to your analyst.",
)
TEMPLATE_EXAMPLE_ROWS: tuple[tuple[object, ...], ...] = (
    ("2024-01-01", 8327258, 1004244, 443069, 342075, 324671, 0, 0, 1.00),
    ("2024-01-08", 8775151, 1354993, 393618, 253124, 342648, 0, 0, 1.00),
    ("2024-01-15", 9101442, 921530, 466911, 398220, 0, 1, 0, 0.98),
    ("2024-01-22", 8894019, 1010877, 421337, 301945, 610500, 0, 0, 1.00),
)


def write_template(path: Path) -> Path:
    """Write the Excel data template (instructions, example data, column descriptions)."""
    workbook = Workbook()
    instructions = workbook.active
    instructions.title = "Instructions"
    for line in TEMPLATE_INSTRUCTIONS:
        instructions.append([line])
    instructions["A1"].font = Font(bold=True, size=14)
    instructions.column_dimensions["A"].width = 110

    data = workbook.create_sheet(config.TEMPLATE_DATA_SHEET)
    data.append(list(TEMPLATE_COLUMNS))
    for row in TEMPLATE_EXAMPLE_ROWS:
        data.append(list(row))
    columns = workbook.create_sheet("Columns")
    columns.append(["Column", "What to put in it"])
    for name, description in TEMPLATE_COLUMNS.items():
        columns.append([name, description])
    for sheet in (data, columns):
        for cell in sheet[1]:
            cell.font = Font(bold=True)
        sheet.freeze_panes = "A2"
        for index in range(1, sheet.max_column + 1):
            sheet.column_dimensions[get_column_letter(index)].width = 22
    columns.column_dimensions["B"].width = 90

    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)
    return path


def read_template(path: Path) -> pd.DataFrame:
    """Read a filled-in template's 'Data' sheet as a contract frame."""
    frame = pd.read_excel(path, sheet_name=config.TEMPLATE_DATA_SHEET)
    frame[config.DATE_COL] = parse_dates(frame[config.DATE_COL])
    return frame
