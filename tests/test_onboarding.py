"""Tests for the export mappers, weekly aggregation, anonymization and the Excel template."""

from pathlib import Path

import pandas as pd
import pytest
from openpyxl import load_workbook

from mixlab import config
from mixlab.onboarding import (
    OnboardingError,
    anonymize,
    build_contract_frame,
    deanonymize,
    factor_from_secret,
    map_generic,
    map_google_ads,
    map_meta_ads,
    map_shopify_orders,
    parse_dates,
    parse_money,
    read_export,
    read_template,
    to_weekly,
    write_template,
)
from mixlab.validate import Severity, validate


def test_parse_money_handles_symbols_grouping_and_negatives() -> None:
    values = pd.Series(["₹1,23,456.50", "Rs 900", "$1,200", "(120)", "", "INR 45.5", "-30"])
    assert parse_money(values).tolist() == [123456.5, 900.0, 1200.0, -120.0, 0.0, 45.5, -30.0]
    assert parse_money(pd.Series([1.5, None])).tolist() == [1.5, 0.0]
    with pytest.raises(OnboardingError, match="amount of money"):
        parse_money(pd.Series(["twelve"]))


def test_parse_dates_reads_iso_with_offsets_and_keeps_the_local_day() -> None:
    parsed = parse_dates(pd.Series(["2024-03-05 23:50:11 +0530", "2024-03-06"]))
    assert parsed.tolist() == [pd.Timestamp("2024-03-05"), pd.Timestamp("2024-03-06")]


def test_parse_dates_infers_day_first_and_month_first() -> None:
    assert parse_dates(pd.Series(["25/03/2024", "01/04/2024"])).tolist() == [
        pd.Timestamp("2024-03-25"),
        pd.Timestamp("2024-04-01"),
    ]
    assert parse_dates(pd.Series(["03/25/2024", "04/01/2024"])).iloc[1] == pd.Timestamp(
        "2024-04-01"
    )
    assert parse_dates(pd.Series(["01/04/2024"]), dayfirst=True).iloc[0] == pd.Timestamp(
        "2024-04-01"
    )
    assert parse_dates(pd.Series(["Mar 5, 2024"])).iloc[0] == pd.Timestamp("2024-03-05")
    with pytest.raises(OnboardingError, match="as a date"):
        parse_dates(pd.Series(["not a date"]))


def test_to_weekly_starts_weeks_on_monday_and_fills_gaps() -> None:
    daily = pd.DataFrame(
        {
            # Sunday 7 Jan belongs to the week of Mon 1 Jan; Mon 22 Jan leaves a gap week.
            "date": pd.to_datetime(["2024-01-01", "2024-01-07", "2024-01-08", "2024-01-22"]),
            "value": [10.0, 5.0, 3.0, 1.0],
        }
    )
    weekly = to_weekly(daily)
    assert weekly["date"].dt.dayofweek.eq(0).all()
    assert weekly["date"].dt.strftime("%Y-%m-%d").tolist() == [
        "2024-01-01",
        "2024-01-08",
        "2024-01-15",
        "2024-01-22",
    ]
    assert weekly["value"].tolist() == [15.0, 3.0, 0.0, 1.0]


def test_meta_export_sums_campaign_rows_per_week() -> None:
    export = pd.DataFrame(
        {
            "Day": ["2024-01-01", "2024-01-01", "2024-01-03", "2024-01-09"],
            "Campaign name": ["Prospecting", "Retargeting", "Prospecting", "Prospecting"],
            "Amount spent (INR)": ["1,000.50", "499.50", "250", "2000"],
        }
    )
    weekly = map_meta_ads(export)
    assert weekly.columns.tolist() == ["date", "spend_meta_ads"]
    assert weekly["spend_meta_ads"].tolist() == [1750.0, 2000.0]


def test_meta_export_without_a_spend_column_explains_itself() -> None:
    with pytest.raises(OnboardingError, match="amount spent"):
        map_meta_ads(pd.DataFrame({"Day": ["2024-01-01"], "Impressions": [10]}))


def test_google_report_with_title_lines_splits_by_campaign_type(tmp_path: Path) -> None:
    path = tmp_path / "google.csv"
    path.write_text(
        "Campaign report\n"
        "1 January 2024 - 14 January 2024\n"
        "Day,Campaign type,Campaign,Cost\n"
        '2024-01-02,Search,Brand,"₹1,500.00"\n'
        "2024-01-03,Video,Launch film,₹700.00\n"
        "2024-01-09,Search,Generic,₹300.00\n"
        "2024-01-10,Demand Gen,Test,₹50.00\n"
        'Total: Campaigns,,,"₹2,550.00"\n'
    )
    weekly = map_google_ads(read_export(path, "Cost"))
    assert set(weekly.columns) == {
        "date",
        "spend_google_search",
        "spend_youtube",
        "spend_google_other",
    }
    assert weekly["spend_google_search"].tolist() == [1500.0, 300.0]
    assert weekly["spend_youtube"].tolist() == [700.0, 0.0]
    assert weekly["spend_google_other"].tolist() == [0.0, 50.0]
    combined = map_google_ads(read_export(path, "Cost"), split_by_campaign_type=False)
    assert combined["spend_google_ads"].tolist() == [2200.0, 350.0]


def shopify_export() -> pd.DataFrame:
    return pd.DataFrame(
        {
            "Name": ["#1001", "#1001", "#1002", "#1003", "#1004", "#1005"],
            "Created at": [
                "2024-01-01 10:00:00 +0530",
                "2024-01-01 10:00:00 +0530",
                "2024-01-07 23:30:00 +0530",
                "2024-01-08 09:00:00 +0530",
                "2024-01-09 09:00:00 +0530",
                "2024-01-10 09:00:00 +0530",
            ],
            "Financial Status": ["paid", None, "partially_refunded", "paid", "refunded", "paid"],
            "Total": [1000.0, None, 500.0, 800.0, 900.0, 700.0],
            "Refunded Amount": [0.0, None, 100.0, 0.0, 900.0, 0.0],
            "Cancelled at": [None, None, None, None, None, "2024-01-10 12:00:00 +0530"],
            "Currency": ["INR"] * 6,
        }
    )


def test_shopify_orders_are_deduplicated_and_net_of_refunds_and_cancellations() -> None:
    weekly = map_shopify_orders(shopify_export())
    # Week 1: order 1001 once (1000) + 1002 net of refund (400). Week 2: only 1003 (800).
    assert weekly["revenue"].tolist() == [1400.0, 800.0]
    assert weekly["date"].dt.strftime("%Y-%m-%d").tolist() == ["2024-01-01", "2024-01-08"]


def test_shopify_rejects_mixed_currencies_and_missing_columns() -> None:
    mixed = shopify_export()
    mixed.loc[3, "Currency"] = "USD"
    with pytest.raises(OnboardingError, match="mixes currencies"):
        map_shopify_orders(mixed)
    with pytest.raises(OnboardingError, match="missing these columns"):
        map_shopify_orders(pd.DataFrame({"Name": ["#1"], "Total": [1.0]}))


def test_contract_frame_joins_on_revenue_weeks_and_zero_fills_spend() -> None:
    revenue = map_shopify_orders(shopify_export())
    meta = pd.DataFrame({"date": pd.to_datetime(["2024-01-08"]), "spend_meta_ads": [250.0]})
    frame = build_contract_frame(revenue, [meta])
    assert frame.columns.tolist() == ["date", "revenue", "spend_meta_ads"]
    assert frame["spend_meta_ads"].tolist() == [0.0, 250.0]
    assert not [i for i in validate(frame).issues if i.check in ("required_columns", "dtypes")]


def test_generic_mapper_produces_a_valid_contract_frame() -> None:
    source = pd.DataFrame(
        {
            "wk": pd.date_range("2022-01-03", periods=110, freq="7D").astype(str),
            "sales": range(1000, 1110),
            "fb": [10.0 + (i % 7) for i in range(110)],
            "tv": [5.0 + (i % 5) for i in range(110)],
            "xmas": [0] * 110,
        }
    )
    frame = map_generic(
        source, "wk", "sales", {"fb": "meta_ads", "tv": "tv"}, {"xmas": "holiday_xmas"}
    )
    assert frame.columns.tolist() == [
        "date",
        "revenue",
        "spend_meta_ads",
        "spend_tv",
        "holiday_xmas",
    ]
    assert validate(frame).count(Severity.CRITICAL) == 0


def test_anonymize_hides_levels_but_preserves_roi_and_is_reversible(df: pd.DataFrame) -> None:
    factor = factor_from_secret("correct horse battery staple")
    scaled = anonymize(df, factor)
    assert scaled["revenue"].iloc[0] == pytest.approx(df["revenue"].iloc[0] * factor)
    ratio_before = df["revenue"].sum() / df["spend_meta_ads"].sum()
    assert scaled["revenue"].sum() / scaled["spend_meta_ads"].sum() == pytest.approx(ratio_before)
    untouched = [config.DATE_COL, config.PROMO_COL, config.PRICE_COL, "holiday_diwali"]
    pd.testing.assert_frame_equal(scaled[untouched], df[untouched])
    pd.testing.assert_frame_equal(deanonymize(scaled, factor), df, check_exact=False)


def test_anonymization_factor_depends_only_on_the_secret() -> None:
    low, high = config.ANONYMIZE_FACTOR_RANGE
    first = factor_from_secret("brand-a")
    assert first == factor_from_secret("brand-a") != factor_from_secret("brand-b")
    assert low <= first < high
    with pytest.raises(OnboardingError, match="cannot be empty"):
        factor_from_secret("")
    with pytest.raises(OnboardingError, match="positive"):
        anonymize(pd.DataFrame({"revenue": [1.0]}), 0.0)


def test_template_has_three_sheets_and_its_data_reads_back(tmp_path: Path) -> None:
    path = write_template(tmp_path / "template.xlsx")
    assert load_workbook(path).sheetnames == ["Instructions", "Data", "Columns"]
    frame = read_template(path)
    assert {"date", "revenue", "spend_meta_ads", "promo_flag", "price_index"} <= set(frame.columns)
    assert frame["date"].dt.dayofweek.eq(0).all()
    assert not [i for i in validate(frame).issues if i.check in ("required_columns", "dtypes")]


def test_shipped_template_exists() -> None:
    assert (config.TEMPLATES_DIR / config.DATA_TEMPLATE_FILENAME).exists()
